from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"D:\桌面\Software Engineering\SE2026-AI-1_需求分析文档_V1.0.docx")

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE = "F3F7FA"
GRAY = "E7E6E6"
BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_table_borders(table, color=BORDER, size='6'):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_run_font(run, east='Microsoft YaHei', west='Arial', size=None, bold=None, color=BLACK):
    run.font.name = west
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), west)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), west)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def set_paragraph_keep(paragraph, keep_next=False, keep_together=False):
    pPr = paragraph._p.get_or_add_pPr()
    if keep_next:
        pPr.append(OxmlElement('w:keepNext'))
    if keep_together:
        pPr.append(OxmlElement('w:keepLines'))

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    set_run_font(run, size=9, color='666666')

def add_para(doc, text='', style=None, bold_lead=None, align=None, before=0, after=6, first_indent=True):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.35
    if first_indent and style is None:
        p.paragraph_format.first_line_indent = Pt(22)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead); set_run_font(r1, bold=True, size=10.5)
        r2 = p.add_run(text[len(bold_lead):]); set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text); set_run_font(r, size=10.5)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text); set_run_font(r, size=10.5)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text); set_run_font(r, size=10.5)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    r = p.add_run(text); set_run_font(r, size=15 if level == 1 else 12, bold=True)
    set_paragraph_keep(p, keep_next=True)
    return p

def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 100, 110, 100, 110)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h)); set_run_font(r, size=9.2, bold=True, color=WHITE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 90, 110, 90, 110)
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val)); set_run_font(r, size=font_size)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.1); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.3)
sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)

# Base styles
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(6)
for name, size, before, after in [('Title', 24, 0, 12), ('Subtitle', 12, 0, 6), ('Heading 1', 15, 16, 8), ('Heading 2', 12, 10, 5), ('Heading 3', 10.5, 8, 4)]:
    st = styles[name]
    st.font.name = 'Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.size = Pt(size); st.font.bold = name.startswith('Heading') or name == 'Title'; st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)

# Footer
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('SE2026-AI-1 需求分析文档  |  '); set_run_font(fr, size=8.5, color='666666')
add_page_number(fp)

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(72); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('SE2026 AI 1'); set_run_font(r, size=16, bold=True, color=NAVY)
p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('基于图像数据的人员关系推断与判别系统'); set_run_font(r, size=24, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(42)
r = p.add_run('需求分析文档'); set_run_font(r, size=18, bold=True)
cover_rows = [
    ['课程', '人工智能班 软件工程课程实践'],
    ['项目编号', 'SE2026-AI-1'],
    ['文档版本', 'V1.0'],
    ['编制日期', '2026年10月19日'],
    ['适用阶段', '需求评审 架构设计 开发测试 验收'],
]
add_table(doc, ['文档信息', '内容'], cover_rows, [4.0, 10.5], 10)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(28); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('项目组成员'); set_run_font(r, size=12, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('沈子妍  谢雨欣  于慧珊  于慧  陈佳仪'); set_run_font(r, size=11)
doc.add_page_break()

add_heading(doc, '文档修订记录', 1)
add_table(doc, ['版本', '日期', '修订内容', '状态'], [['V1.0', '2026-10-19', '根据课程选题要求和团队分工形成首版需求基线', '待团队评审']], [2,3,8,3])
add_heading(doc, '项目成员与职责', 1)
team_rows = [
    ['沈子妍', '产品经理', '组织需求调研与范围管理；维护需求基线和优先级；协调评审、演示和验收；确认用户流程与产品交付物。'],
    ['谢雨欣', '架构师', '设计总体架构、数据流和接口边界；确定图像分析、图推理、图数据库与前后端集成方案；负责关键技术风险评审。'],
    ['于慧珊', '运维工程师', '设计开发测试环境、容器化部署、配置与日志方案；负责服务监控、备份恢复、发布回滚及数据存储合规检查。'],
    ['于慧', '开发工程师', '实现核心后端和算法集成；完成检测跟踪、ReID、关系推断、图谱查询接口及前端功能联调。'],
    ['陈佳仪', '测试工程师', '编制测试计划与用例；负责功能、接口、性能、算法指标、安全与回归测试；维护缺陷与验收证据。'],
]
add_table(doc, ['成员', '岗位', '主要职责'], team_rows, [2.6,3,10])
add_heading(doc, '文档目的', 1)
add_para(doc, '本文档定义系统在课程项目范围内必须具备的业务能力、功能、质量属性、数据与合规约束以及验收条件。团队后续的架构设计、开发任务、测试用例和演示脚本均应能追溯到本需求基线。')
add_para(doc, '当前版本以课程选题页和项目分工为依据。选题页未给出明确阈值的算法指标将在获得数据集并完成基线实验后冻结；文中标注为待确认的事项不应被视为已承诺的验收数值。')
add_heading(doc, '目录', 1)
toc_items = ['1 项目概述', '2 用户与使用场景', '3 系统范围', '4 功能需求', '5 数据需求', '6 算法与模型需求', '7 接口与交互需求', '8 非功能需求', '9 安全隐私与伦理要求', '10 部署运维需求', '11 测试与验收', '12 项目约束与风险', '13 需求优先级与迭代计划', '14 待确认事项', '附录 A 需求追踪矩阵']
for item in toc_items: add_para(doc, item, first_indent=False, after=2)
doc.add_page_break()

add_heading(doc, '1 项目概述', 1)
add_heading(doc, '1.1 项目背景', 2)
add_para(doc, '本项目面向图像和视频数据中的人员关系分析。系统从图像或视频中检测并跟踪人员，提取可用于跨镜头匹配的行人特征，结合共同出现、时空接近和交互频次等信息构建人员关系图，再利用图神经网络等方法推断关系强度或关系类别。系统以可视化图谱和查询结果向用户呈现推断依据与结果。')
add_para(doc, '需求依据包括课程实践选题页中的 SE2026-AI-01 课题说明及项目组岗位分工。课程明确要求关系强度计算、人员关系图谱构建与查询、关系判别可视化，并涉及图数据库建模、隐私计算和单元测试。本文以项目组使用的 SE2026-AI-1 作为项目编号。除课程明确条目外，本文细化的角色、接口、异常策略和迭代范围均作为团队评审方案，评审确认后形成执行基线。')
add_heading(doc, '1.2 建设目标', 2)
for t in [
    '建立从图像或视频输入、人员检测跟踪、ReID 特征提取、关系特征生成、图推断、图数据库存储到可视化查询的完整工程链路。',
    '支持基于节点和边特征的关系强度计算，并能够形成、更新和查询人员关系图谱。',
    '对关系判别结果提供可理解的可视化信息，包括人员节点、关系类型或强度、置信度、时间范围及可用的证据摘要。',
    '落实匿名化、最小权限、访问审计和合规存储要求，避免系统被用于未授权的身份识别或监控。',
    '形成可测试、可部署、可演示的软件系统及配套文档。'
]: add_bullet(doc, t)
add_heading(doc, '1.3 成功标准', 2)
success_rows = [
    ['端到端流程', '能够导入合规数据并完成检测、跟踪、特征生成、关系推断、入图和展示。', '演示脚本及系统日志'],
    ['ReID 能力', '输出 mAP、CMC 等约定指标；具体达标阈值在基线实验后由团队评审冻结。', '算法评测报告'],
    ['关系推断', '输出准确率、宏平均 F1 等指标；具体阈值在标签方案与数据划分确认后冻结。', '独立测试集报告'],
    ['图查询性能', '课程规定的典型图查询响应时间小于 1 秒。', '性能测试记录'],
    ['隐私合规', '默认不向普通用户暴露真实身份信息，关键访问均可审计。', '安全测试及审计记录'],
]
add_table(doc, ['维度', '成功标准', '验证方式'], success_rows, [3,8,5])

add_heading(doc, '2 用户与使用场景', 1)
add_heading(doc, '2.1 用户角色', 2)
roles = [
    ['系统管理员', '维护用户、角色、数据源、模型版本和系统配置；查看审计日志与运行状态。'],
    ['分析人员', '上传或选择数据任务，发起分析，查询关系图谱，查看关系判别及证据摘要，导出脱敏结果。'],
    ['算法开发人员', '管理模型与评测配置，触发离线评测，查看模型指标和错误样本，不直接获得无授权的真实身份信息。'],
    ['审计人员', '只读查看数据处理记录、访问日志、模型版本、导出行为和合规配置。'],
]
add_table(doc, ['角色', '权限与职责'], roles, [4,12])
add_heading(doc, '2.2 核心使用场景', 2)
scenarios = [
    ['UC01 创建分析任务', '分析人员选择图像或视频数据，设置时间范围和模型版本，提交异步分析任务。'],
    ['UC02 查看处理进度', '用户查看任务排队、执行、完成或失败状态，并获得可定位的错误信息。'],
    ['UC03 浏览关系图谱', '用户按人员匿名标识、时间或关系类型浏览节点与边，缩放、筛选和查看详情。'],
    ['UC04 查询两人关系', '用户选择两个匿名人员节点，查询关系强度、预测类别、置信度和支持证据摘要。'],
    ['UC05 追溯推断结果', '用户查看该结果使用的数据时间段、模型版本、主要特征和生成时间。'],
    ['UC06 管理与审计', '管理员配置权限、数据保留期和服务参数；审计人员查询访问和导出记录。'],
]
add_table(doc, ['编号', '场景说明'], scenarios, [3,13])

add_heading(doc, '3 系统范围', 1)
add_heading(doc, '3.1 范围内能力', 2)
for x in ['图像及视频文件导入和数据校验。', '人员检测、跨帧跟踪、ReID 特征提取和匿名人员标识生成。', '共同出现、相对距离、持续时间、交互频次等关系特征计算。', '人员关系强度计算及关系类别推断。', 'Neo4j 人员关系图谱建模、写入、更新和查询。', 'Vue 3 可视化界面及 FastAPI 服务接口。', '用户权限、操作审计、匿名化与数据生命周期管理。', '模型和系统的功能、性能及算法评测。']: add_bullet(doc, x)
add_heading(doc, '3.2 范围外能力', 2)
for x in ['将系统用于公共场所实时大规模监控或执法决策。', '仅凭系统推断结果对个人作出惩罚、信用评价或其他高风险自动化决策。', '在未取得合法授权时将匿名人员标识还原为真实身份。', '保证推断结果代表现实中的法律、亲属、雇佣或情感关系。', '首版课程项目中的移动端原生应用和超大规模分布式训练平台。']: add_bullet(doc, x)
add_heading(doc, '3.3 系统边界与上下文', 2)
add_para(doc, '输入端接收经授权的图像、视频及可选关系标签；算法服务完成检测、跟踪、特征提取和推断；图数据库保存匿名节点、关系边和必要的可追溯元数据；后端 API 对外提供任务、查询、模型和审计能力；Web 前端负责管理操作和图谱展示。原始数据存储、模型文件与日志存储均属于受控资源，不允许由前端直接访问。')

add_heading(doc, '4 功能需求', 1)
functional = [
    ['FR-01', '用户认证与授权', '系统应支持登录、退出和基于角色的访问控制；未授权请求应被拒绝并记录。', '必须'],
    ['FR-02', '数据导入', '系统应支持常用图像格式和课程约定的视频格式，校验类型、大小、完整性和授权信息。', '必须'],
    ['FR-03', '任务管理', '系统应支持创建、启动、取消、查询和重试分析任务，并保留状态、配置、模型版本和时间记录。', '必须'],
    ['FR-04', '人员检测', '系统应使用 YOLOv8 或经评审的替代模型识别人员目标，输出边界框、置信度和帧信息。', '必须'],
    ['FR-05', '人员跟踪', '系统应使用 ByteTrack 或经评审的替代方案关联连续帧中的人员轨迹，并处理短时遮挡。', '必须'],
    ['FR-06', 'ReID 特征', '系统应生成匿名人员特征和跨镜头匹配结果，保存模型版本及相似度，不保存不必要的可逆身份信息。', '必须'],
    ['FR-07', '关系特征', '系统应按配置窗口计算共同出现次数、持续时间、时空距离、运动一致性等关系特征。', '必须'],
    ['FR-08', '关系强度', '系统应依据节点及边特征计算归一化关系强度，并记录算法版本和计算时间。', '必须'],
    ['FR-09', '关系判别', '系统应利用 GNN 或基线模型输出关系类别或关系存在性、置信度及模型版本。', '必须'],
    ['FR-10', '图谱构建', '系统应在 Neo4j 中创建或更新匿名人员节点、关系边和分析任务关联，避免重复节点和失效引用。', '必须'],
    ['FR-11', '图谱查询', '系统应支持按匿名人员、时间范围、关系强度、关系类别和任务查询节点、边及邻居。', '必须'],
    ['FR-12', '结果可视化', '系统应以关系图展示节点与边，支持缩放、拖动、筛选、详情查看和置信度提示。', '必须'],
    ['FR-13', '结果追溯', '每个推断结果应能追溯到任务、数据时间范围、模型版本和主要特征摘要。', '必须'],
    ['FR-14', '结果导出', '系统可导出经权限校验和脱敏的查询结果；导出字段、时间、操作者和用途应记录。', '应该'],
    ['FR-15', '模型管理', '管理员应能登记模型名称、版本、指标、状态和文件校验值，并指定当前可用版本。', '应该'],
    ['FR-16', '审计查询', '审计角色应能按用户、时间、操作类型和对象查询访问、配置变更及导出日志。', '必须'],
    ['FR-17', '异常处理', '系统应对输入错误、模型失败、数据库不可用和超时给出明确状态，不得生成伪造的成功结果。', '必须'],
]
add_table(doc, ['编号', '名称', '需求描述', '优先级'], functional, [2.1,2.8,9.2,2], 8.7)

add_heading(doc, '5 数据需求', 1)
add_heading(doc, '5.1 数据对象', 2)
data_rows = [
    ['媒体数据', '图像或视频、采集时间、来源、授权状态、校验值、保留期限', '原始数据与公开展示隔离'],
    ['检测记录', '任务、帧号、边界框、置信度、检测模型版本', '可按保留策略清理'],
    ['轨迹记录', '匿名轨迹标识、帧区间、位置序列、质量分数', '不得直接等同真实身份'],
    ['人员节点', '匿名人员标识、特征引用、首次和末次出现、节点版本', '前端只显示匿名标识'],
    ['关系边', '源节点、目标节点、关系类别、强度、置信度、有效时间', '方向性由关系定义决定'],
    ['推断证据', '共同出现、持续时间、距离等聚合特征及模型版本', '避免暴露原始敏感画面'],
    ['审计日志', '操作者、操作、对象、时间、结果、来源地址', '只追加并限制访问'],
]
add_table(doc, ['数据对象', '主要字段', '管理要求'], data_rows, [3,8,5])
add_heading(doc, '5.2 数据质量要求', 2)
for x in ['输入数据必须具有来源和使用授权记录；无法确认授权的数据不得进入正式分析环境。', '训练集、验证集和测试集应按人员身份或视频源进行隔离，避免同一对象泄漏导致指标虚高。', '标注规范应定义关系类别、无关系样本、模糊样本及标注冲突处理，并统计一致性。', '媒体、检测、轨迹、节点和关系记录之间应通过稳定标识关联，孤立引用和重复键应可检测。', '数据清洗、标注修订和特征生成应保留版本，支持复现实验。']: add_bullet(doc, x)
add_heading(doc, '5.3 图数据库模型', 2)
add_para(doc, '建议的核心节点包括 PersonAnon、Track、Media、AnalysisTask 和 ModelVersion；核心关系包括 APPEARS_IN、MATCHED_AS、CO_OCCURS_WITH 和 INFERRED_RELATION。PersonAnon 仅表示系统生成的匿名实体，不等同于现实身份。关系边应包含强度、类别、置信度、有效时间、来源任务和模型版本，以便查询和追溯。')

add_heading(doc, '6 算法与模型需求', 1)
algo = [
    ['AR-01', '检测与跟踪', '建立 YOLOv8 与 ByteTrack 基线；对遮挡、漏检、轨迹切换保存质量统计。'],
    ['AR-02', 'ReID 评测', '使用固定的数据划分和评测脚本报告 mAP、CMC Rank-1 等指标，并保存随机种子与模型校验值。'],
    ['AR-03', '特征规范化', '相似度、距离、持续时间和频次等输入应具有明确的归一化方法，训练与推理保持一致。'],
    ['AR-04', '图构建规则', '定义节点合并、边生成、时间窗口和负样本策略，防止由同一轨迹产生自关系。'],
    ['AR-05', '关系模型', '使用 PyTorch Geometric 实现 GNN 或经评审的图模型，并保留非图基线用于对照。'],
    ['AR-06', '输出约束', '输出必须包含类别或强度、置信度、模型版本和适用时间，不得只给出不可追溯结论。'],
    ['AR-07', '不确定性', '低置信度或证据不足的结果应标记为不确定，不得强制归类。'],
    ['AR-08', '可复现性', '训练配置、依赖版本、数据版本、随机种子和指标文件应纳入版本管理或实验记录。'],
]
add_table(doc, ['编号', '主题', '要求'], algo, [2.2,3.2,10.6])
add_heading(doc, '6.1 关系定义原则', 2)
add_para(doc, '系统推断的是图像数据支持的统计关联或预定义互动类型，不应自动解释为亲属、朋友、同事等现实社会关系。若课程演示需要使用具体关系标签，产品经理和测试工程师应先确认标签定义、证据标准和误用提示，并在界面中显示免责声明。')
add_heading(doc, '6.2 指标冻结机制', 2)
add_para(doc, '课程要求明确关注 ReID mAP、关系推断准确率和图查询响应时间。图查询响应时间以小于 1 秒为验收目标。ReID mAP 和关系推断准确率的具体阈值应在数据集、划分方式、标签体系和基线模型确定后冻结，至少同时报告宏平均 F1、混淆矩阵和样本规模，避免单一准确率掩盖类别不平衡。')

add_heading(doc, '7 接口与交互需求', 1)
add_heading(doc, '7.1 后端接口', 2)
apis = [
    ['认证', '/api/auth/*', '登录、退出、当前用户及令牌刷新'],
    ['任务', '/api/tasks/*', '创建任务、查询进度、取消、重试和结果摘要'],
    ['媒体', '/api/media/*', '受控上传、元数据查询及删除申请'],
    ['图谱', '/api/graph/*', '节点查询、邻居查询、路径查询和关系过滤'],
    ['推断', '/api/inference/*', '两人关系查询、证据摘要和模型信息'],
    ['模型', '/api/models/*', '模型登记、指标查询和状态切换'],
    ['审计', '/api/audit/*', '按权限查询审计日志和导出记录'],
]
add_table(doc, ['模块', '建议路径', '主要能力'], apis, [3,4.5,8.5])
add_para(doc, '接口采用 JSON 作为主要数据格式，上传接口使用受限的 multipart/form-data。所有响应应包含请求标识，错误响应包含稳定错误码和可理解说明。耗时分析任务使用异步任务模式，前端通过轮询或服务端事件获取进度，不应让单次 HTTP 请求长期占用。')
add_heading(doc, '7.2 前端交互', 2)
for x in ['首页显示任务概览、近期任务和服务健康摘要。', '任务创建页明确显示支持的格式、隐私提示、模型版本和参数默认值。', '图谱页支持节点搜索、关系筛选、时间筛选、布局切换、图例和详情面板。', '置信度、关系强度和预测类别使用不同视觉编码，并提供文字说明，不能仅依赖颜色。', '失败和不确定结果应在界面中清楚区分；不得以空图或默认值冒充成功。', '用户执行导出、删除、模型切换等敏感操作前应获得确认并显示结果。']: add_bullet(doc, x)

add_heading(doc, '8 非功能需求', 1)
nfr = [
    ['NFR-01', '性能', '在课程规定的测试环境和约定数据规模下，典型图查询的服务端响应时间应小于 1 秒；批量分析任务应显示进度。'],
    ['NFR-02', '可靠性', '任务状态应持久化；服务重启后已完成结果可查询，失败任务可定位并按策略重试。'],
    ['NFR-03', '可用性', '常见操作应有清晰反馈、空状态和错误提示；核心演示流程不依赖命令行。'],
    ['NFR-04', '可维护性', '前端、后端、算法和数据访问模块边界清晰；关键模块具有单元测试和接口文档。'],
    ['NFR-05', '可扩展性', '检测、ReID 和关系模型通过统一接口封装，允许替换模型而不重写业务流程。'],
    ['NFR-06', '兼容性', 'Web 界面至少支持课程演示所用的主流 Chromium 浏览器；后端环境由容器或锁定依赖保证一致。'],
    ['NFR-07', '可观测性', '服务输出结构化日志和请求标识；记录任务耗时、失败次数、队列长度及数据库健康状态。'],
    ['NFR-08', '可测试性', '核心计算提供确定性测试输入；外部依赖可模拟；测试数据与生产或真实数据隔离。'],
]
add_table(doc, ['编号', '类别', '要求'], nfr, [2.2,2.6,11.2])

add_heading(doc, '9 安全隐私与伦理要求', 1)
security = [
    ['SR-01', '数据最小化', '仅收集完成功能所需的媒体、特征和元数据；默认不保存多余的人脸或身份信息。'],
    ['SR-02', '匿名化', '人员节点使用不可读匿名标识；真实身份映射如确需存在，应独立加密存储并严格授权。'],
    ['SR-03', '访问控制', '采用最小权限和角色授权；管理、模型切换、导出和删除操作仅向获授权角色开放。'],
    ['SR-04', '传输与存储', '敏感数据传输使用加密通道；密钥、口令和令牌不得写入代码仓库或普通日志。'],
    ['SR-05', '访问审计', '登录、数据访问、查询、导出、配置变化、删除和授权失败均应记录审计日志。'],
    ['SR-06', '数据生命周期', '为原始媒体、派生特征、图谱结果和日志配置保留期限及可验证删除流程。'],
    ['SR-07', '结果限制', '界面应声明推断存在误差，结果不得作为高风险决策的唯一依据。'],
    ['SR-08', '数据合规', '仅使用课程允许、公开许可或取得授权的数据集，并遵守其许可证与使用限制。'],
    ['SR-09', '联邦或脱敏研究', '若实施联邦学习，应记录参与方、参数交换和攻击面；若不实施，应至少完成脱敏计算流程。'],
]
add_table(doc, ['编号', '主题', '要求'], security, [2.2,3,10.8])

add_heading(doc, '10 部署运维需求', 1)
ops = [
    ['环境', '区分开发、测试和演示环境；配置通过环境变量或受控配置文件注入。'],
    ['部署', 'FastAPI、前端、Neo4j 和算法服务使用可复现的容器或锁定环境部署，并提供启动说明。'],
    ['健康检查', '提供 API、图数据库及必要算法服务的健康检查，前端可展示不可用状态。'],
    ['日志', '应用、访问、任务和审计日志分级记录；不得记录原始密钥、完整令牌或不必要的敏感特征。'],
    ['监控', '至少监控服务存活、请求错误、任务积压、处理耗时、磁盘使用和数据库连接。'],
    ['备份恢复', '备份图数据库配置和必要结果；在演示前执行一次恢复验证并记录恢复步骤。'],
    ['发布回滚', '发布包具有版本号；数据库结构和模型版本变更前备份，失败时可回滚至上一稳定版本。'],
]
add_table(doc, ['领域', '要求'], ops, [4,12])

add_heading(doc, '11 测试与验收', 1)
add_heading(doc, '11.1 测试范围', 2)
test_rows = [
    ['单元测试', '关系特征计算、强度计算、数据校验、权限判断、图数据库访问封装和错误处理。'],
    ['接口测试', '认证、任务、图谱、推断、模型及审计 API 的正常、边界和未授权场景。'],
    ['算法测试', '固定测试集上的检测、ReID 和关系推断指标，可复现性及错误案例分析。'],
    ['集成测试', '媒体导入至图谱展示的完整链路，以及模型、API、Neo4j 与前端之间的数据一致性。'],
    ['性能测试', '典型图查询小于 1 秒；并记录并发、数据规模、硬件环境、P50、P95 和最大耗时。'],
    ['安全测试', '越权、弱口令策略、敏感日志、未授权导出、文件类型绕过、审计完整性。'],
    ['恢复测试', '服务重启、数据库暂时不可用和任务失败后的状态恢复与错误提示。'],
    ['验收演示', '按预先冻结脚本完成数据导入、任务执行、图谱浏览、关系查询、审计查看和异常展示。'],
]
add_table(doc, ['测试类型', '重点内容'], test_rows, [4,12])
add_heading(doc, '11.2 核心验收条件', 2)
for x in ['所有必须级功能需求均通过测试，且不存在阻断演示的严重缺陷。', '端到端分析流程在指定演示数据上成功完成，结果可在图谱界面查询并追溯。', 'ReID 和关系推断报告包含冻结的数据划分、指标、模型版本、样本数和错误分析。', '典型图查询在约定环境和规模下响应时间小于 1 秒，并提供原始测试记录。', '普通分析人员无法查看管理接口、完整审计日志或未授权数据。', '系统不会将低置信度结果显示为确定事实，导出数据保持脱敏。', '部署文档、接口文档、测试报告、用户说明和演示脚本齐全。']: add_bullet(doc, x)
add_heading(doc, '11.3 需求完成定义', 2)
add_para(doc, '一项需求只有在实现代码合并、自动或手工测试通过、必要文档更新、产品经理确认交互、测试工程师保留证据且未引入严重回归后，才可标记为完成。涉及模型的需求还必须绑定模型版本与评测结果；涉及数据的需求还必须通过授权和脱敏检查。')

add_heading(doc, '12 项目约束与风险', 1)
risks = [
    ['数据与标签不足', '关系标签主观、类别不平衡或公开数据缺少真实关系标签。', '先定义可观察的关系任务；设置基线；使用宏平均 F1；记录标注一致性。', '产品经理 测试工程师'],
    ['ReID 错配传播', '身份匹配错误会造成错误节点合并和关系边。', '设置质量阈值与不确定状态；保留轨迹层；进行错误样本审查。', '架构师 开发工程师'],
    ['隐私风险', '图像和人员特征具有敏感性，导出可能造成重新识别。', '匿名化、最小权限、限制导出、审计、保留期和演示数据隔离。', '产品经理 运维工程师'],
    ['性能不稳定', '图查询随数据规模和查询深度增长而变慢。', '限制默认查询深度；建立索引；分页；对典型查询做性能基线。', '架构师 运维工程师'],
    ['算力限制', '课程环境可能无法实时完成全部模型推理。', '采用异步任务和离线预计算；提供轻量模型或演示缓存，但必须标识结果来源。', '架构师 开发工程师'],
    ['技术集成复杂', 'YOLOv8、ByteTrack、PyG、Neo4j、FastAPI 和 Vue3 的数据结构不一致。', '先冻结接口契约与最小贯通样例，再并行扩展模块。', '架构师'],
    ['结果误用', '用户可能把统计关联解释为真实社会关系。', '限制标签语义；显示置信度、证据与免责声明；禁止高风险自动决策。', '产品经理'],
]
add_table(doc, ['风险', '影响', '应对措施', '责任角色'], risks, [2.5,4,6.5,3], 8.5)

add_heading(doc, '13 需求优先级与迭代计划', 1)
plan = [
    ['迭代一 最小贯通', '数据导入、YOLOv8 检测、ByteTrack 跟踪、基础 ReID、任务 API、Neo4j 最小模型、基础图谱页面。', '端到端样例可运行并展示匿名节点。'],
    ['迭代二 关系推断', '关系特征、强度计算、基线关系模型、PyG 图模型、结果追溯和离线评测。', '能够查询两人关系并生成指标报告。'],
    ['迭代三 工程完善', '权限、审计、脱敏导出、异常恢复、性能优化、单元与集成测试、部署监控。', '满足验收条件并形成演示版本。'],
]
add_table(doc, ['阶段', '范围', '退出条件'], plan, [3.4,8.2,4.4])
add_para(doc, '优先级采用必须、应该、可以三级。必须项构成课程验收的最小范围；应该项在不影响核心链路和合规要求的前提下实现；可以项仅在时间和数据条件允许时纳入。团队不得以可视化美化替代核心算法评测、权限或审计要求。')

add_heading(doc, '14 待确认事项', 1)
pending = [
    ['Q-01', '课程最终验收使用的数据集、许可和数据规模。', '产品经理', '需求评审后'],
    ['Q-02', '关系任务是二分类、多分类还是连续强度预测，以及每类的可观察定义。', '产品经理 架构师', '模型设计前'],
    ['Q-03', 'ReID mAP 和关系推断指标的具体验收阈值。', '全体评审', '基线实验后'],
    ['Q-04', '图查询小于 1 秒的查询集合、并发数、数据规模和计时边界。', '架构师 测试工程师', '性能测试前'],
    ['Q-05', '是否要求真正实现联邦学习，或以脱敏计算与设计说明满足隐私工程要求。', '产品经理 任课教师', '迭代一结束前'],
    ['Q-06', '课程演示部署环境的 GPU、内存、操作系统和网络限制。', '运维工程师', '部署设计前'],
    ['Q-07', '是否允许展示原始图像证据；若允许，哪些角色和场景可见。', '产品经理', '前端详情页设计前'],
]
add_table(doc, ['编号', '待确认内容', '责任人', '截止点'], pending, [2.2,8.2,3.2,2.4])

add_heading(doc, '附录 A 需求追踪矩阵', 1)
trace = [
    ['业务目标一 完整分析链路', 'FR-02 至 FR-13', 'AR-01 至 AR-08', '单元 接口 集成 算法 验收演示'],
    ['业务目标二 图谱构建与查询', 'FR-10 至 FR-13', 'NFR-01 NFR-07', '接口 集成 性能'],
    ['业务目标三 隐私合规', 'FR-01 FR-14 FR-16', 'SR-01 至 SR-09', '安全 审计 验收演示'],
    ['业务目标四 工程可交付', 'FR-03 FR-15 FR-17', 'NFR-02 至 NFR-08', '单元 集成 恢复 部署验证'],
]
add_table(doc, ['目标', '功能需求', '质量与算法需求', '验证活动'], trace, [4.2,4.2,4.2,3.4])
add_heading(doc, '附录 B 术语说明', 1)
terms = [
    ['ReID', 'Person Re Identification，行人重识别，用于判断不同图像或镜头中的人员是否可能为同一匿名实体。'],
    ['mAP', 'Mean Average Precision，检索任务的平均精度均值。'],
    ['CMC', 'Cumulative Matching Characteristics，累计匹配特性，常报告 Rank 1 等指标。'],
    ['GNN', 'Graph Neural Network，图神经网络，用节点、边及邻域信息进行学习和推断。'],
    ['关系强度', '由共同出现、时空距离、持续时间等可观察特征计算的归一化关联程度，不代表现实关系事实。'],
    ['匿名人员标识', '系统为数据处理生成的内部标识，不应直接包含姓名、证件号等真实身份信息。'],
]
add_table(doc, ['术语', '说明'], terms, [4,12])

add_heading(doc, '附录 C 核心用例详细说明', 1)
add_heading(doc, 'C1 分析任务与关系图谱生成', 2)
add_para(doc, '参与者为分析人员。前置条件是用户已登录且具有数据使用权限，媒体授权信息完整，系统存在可用模型版本。成功后应生成可追溯的任务记录和关系图谱；失败时应保留失败阶段与原因。')
for x in ['用户上传媒体并选择时间范围、模型和分析参数，系统完成类型、大小、授权和参数校验。', '系统创建任务并返回唯一任务标识，用户可查看排队和执行状态。', '后台依次完成检测跟踪、ReID、关系特征提取和关系推断，保存各阶段状态。', '系统将匿名节点、关系及来源元数据写入图数据库，任务标记完成。', '用户进入图谱页查看结果；若无有效人员或关系，应显示成功处理但无可展示结果，并说明原因。']: add_number(doc, x)
add_para(doc, '异常分支包括无权限、文件损坏、无人员、资源不足、模型加载失败和数据库写入失败。写入应以任务或批次保持一致性；未完成的部分结果不得作为最终成功结果展示。重试应具有幂等控制，避免重复节点和边。取消任务后，应停止后续阶段并标记已保留的中间数据。')
add_heading(doc, 'C2 两人关系查询与人工复核', 2)
add_para(doc, '参与者为分析人员。前置条件是两个匿名节点存在且用户有权访问其所属任务。用户选择节点及时间范围，系统校验权限后返回关系强度、类别、置信度、证据摘要和模型版本。结果为不确定时，应提示证据不足；不存在关系记录时，应显示未检索到记录，不能自动等同于现实中无关系。用户可回到关联任务复核结果，并按授权导出脱敏记录。')

# Avoid orphaned headings and set core properties
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        set_paragraph_keep(p, keep_next=True)
doc.core_properties.title = 'SE2026-AI-1 基于图像数据的人员关系推断与判别系统需求分析文档'
doc.core_properties.subject = '软件工程课程实践需求分析'
doc.core_properties.author = 'SE2026-AI-1 项目组'
doc.core_properties.keywords = '需求分析, ReID, GNN, Neo4j, FastAPI, Vue3'

doc.save(OUT)
print(OUT)
